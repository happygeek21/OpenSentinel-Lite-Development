"""Generate formatted submission documents (DOCX, HTML) and bundle the Google Drive submission folder.
"""

import os
import shutil
import zipfile
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

PROJECT_ROOT = Path(__file__).resolve().parent
SUBMISSION_DIR = Path("C:/Users/Hemanth/Desktop/Cybersecurity_Assessment_Submission")


def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)


def build_docx(output_path: Path):
    doc = docx.Document()

    # Set normal style
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    # Title
    title = doc.add_paragraph()
    r_title = title.add_run("Cybersecurity Assessment Submission Report")
    r_title.bold = True
    r_title.font.size = Pt(24)
    r_title.font.color.rgb = RGBColor(0x1a, 0x36, 0x5d) # Navy
    title.paragraph_format.space_after = Pt(4)

    # Subtitle / Candidate Info
    meta = doc.add_paragraph()
    meta.add_run("Candidate Name: ").bold = True
    meta.add_run("Hemanth Prabhu\n")
    meta.add_run("Candidate Email: ").bold = True
    meta.add_run("hemanthprabhu2003@gmail.com\n")
    meta.add_run("Project Name: ").bold = True
    meta.add_run("OpenSentinel Lite\n")
    meta.add_run("Selected Tasks: ").bold = True
    meta.add_run("Task 13 (Build a Defensive Tool) & Task 9 (Test a Detection Rule)\n")
    meta.add_run("GitHub Repository: ").bold = True
    meta.add_run("https://github.com/happygeek21/OpenSentinel-Lite-Development\n")
    meta.paragraph_format.space_after = Pt(14)

    doc.add_heading("1. Executive Summary", level=1)
    p = doc.add_paragraph(
        "This submission fulfills the requirements of the Cybersecurity Assessment Task Catalogue by completing "
        "two complementary problem statements: Task 13 (Primary) and Task 9 (Secondary).\n\n"
        "1. Task 13 (Defensive Security Tool): Designed, built, and tested OpenSentinel Lite, an automated defensive "
        "monitoring and threat detection system featuring log parsing, heuristic classification (SQL Injection, "
        "Reconnaissance, and Brute Force), dynamic risk scoring, SQLite persistence, a FastAPI service layer, and a real-time React dashboard.\n\n"
        "2. Task 9 (Detection Rule Testing): Empirically evaluated the baseline detection heuristics against a synthetic test corpus "
        "of attack and benign payloads. We identified a 50% false positive rate in baseline substring rules and engineered an "
        "improved regular expression engine with word boundaries that achieved 100% recall with 0% false positives."
    )

    doc.add_heading("2. Task 13: Build a Small Defensive Security Tool (OpenSentinel Lite)", level=1)
    doc.add_heading("2.1 Problem Statement & Objectives", level=2)
    doc.add_paragraph(
        "Modern systems require lightweight, automated mechanisms to detect intrusions and operational anomalies without the overhead "
        "of enterprise SIEM clusters. The objective was to create a modular defensive tool capable of ingesting security logs, "
        "normalizing data, applying real-time threat detection, persisting alerts with calculated severity, and surfacing actionable telemetry."
    )

    doc.add_heading("2.2 System Architecture & Components", level=2)
    doc.add_paragraph(
        "• Log Ingestion & Parsing (backend/parser.py): Normalizes pipe-delimited security logs into structured event dictionaries.\n"
        "• Threat Analyzer Engine (backend/analyzer.py): Detects SQL injection, directory traversal / recon, and multi-attempt brute-force attacks.\n"
        "• Database & Alert Storage (backend/database.py): Relational SQLite schema storing events, alerts, notifications, and blocked IPs.\n"
        "• Service & Visualization Layer (backend/main.py + frontend/): FastAPI REST service and Vite+React interactive dashboard."
    )

    doc.add_heading("2.3 Sample Output Data", level=2)
    p_code = doc.add_paragraph()
    r = p_code.add_run(
        "Sample Log Event:\n"
        "2026-10-29T10:58:09 | SEARCH | user=employee_demo | ip=127.0.0.1 | status=REQUEST | details=Search query: ' OR 1=1 --\n\n"
        "Stored Threat Alert:\n"
        "{\n"
        '  "type": "SQL_INJECTION",\n'
        '  "ip": "127.0.0.1",\n'
        '  "username": "employee_demo",\n'
        '  "attempts": 2,\n'
        '  "risk_score": 90,\n'
        '  "severity": "HIGH",\n'
        '  "detection_method": "SQL Injection Rule"\n'
        "}"
    )
    r.font.name = 'Consolas'
    r.font.size = Pt(9.5)

    doc.add_heading("3. Task 9: Test a Defensive Detection Rule", level=1)
    doc.add_heading("3.1 Hypothesis & Baseline Analysis", level=2)
    doc.add_paragraph(
        "Detection rules relying on naive substring keywords ('select ', 'update ', 'drop table') cause high false positive rates "
        "and alert fatigue because normal users frequently type these words in benign business contexts."
    )

    doc.add_heading("3.2 Quantitative Benchmark Results", level=2)

    # Table
    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    headers = ["Test Payload", "Category", "Baseline", "Improved", "Outcome"]
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        hdr_cells[i].paragraphs[0].runs[0].bold = True
        set_cell_background(hdr_cells[i], "1A365D")
        hdr_cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    data = [
        ("' OR 1=1 --", "Attack", "Flagged", "Flagged", "True Positive"),
        ('" OR 1=1 #', "Attack", "Flagged", "Flagged", "True Positive"),
        ("admin' --", "Attack", "Flagged", "Flagged", "True Positive"),
        ("1' UNION SELECT ...", "Attack", "Flagged", "Flagged", "True Positive"),
        ("1; DROP TABLE logs; --", "Attack", "Flagged", "Flagged", "True Positive"),
        ("test') UNION ALL SELECT ...", "Attack", "Flagged", "Flagged", "True Positive"),
        ("Please select insurance plan", "Benign", "FLAGGED (FP)", "Clean", "False Alarm Eliminated"),
        ("Need to update contact info", "Benign", "FLAGGED (FP)", "Clean", "False Alarm Eliminated"),
        ("Where can I drop off files?", "Benign", "Clean", "Clean", "True Negative"),
        ("Review of select committee report", "Benign", "FLAGGED (FP)", "Clean", "False Alarm Eliminated"),
        ("Order 101 or 102 was delivered", "Benign", "Clean", "Clean", "True Negative"),
        ("Employee salary directory", "Benign", "Clean", "Clean", "True Negative"),
    ]

    for row in data:
        row_cells = table.add_row().cells
        for i, val in enumerate(row):
            row_cells[i].text = val
            if "FP" in val:
                row_cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
            elif "Eliminated" in val or "True Negative" in val or "True Positive" in val:
                row_cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(0x00, 0x70, 0x20)

    doc.add_paragraph("\nPerformance Summary:")
    doc.add_paragraph(
        "• Baseline Rule: 100% Recall (6/6 attacks caught), 50% False Positive Rate (3/6 benign queries incorrectly flagged).\n"
        "• Improved Rule: 100% Recall (6/6 attacks caught), 0% False Positive Rate (0/6 benign queries flagged)."
    )

    doc.add_heading("4. Reproduction Guide & Links", level=1)
    doc.add_paragraph(
        "• GitHub Repository: https://github.com/happygeek21/OpenSentinel-Lite-Development\n"
        "• Unit Test Suite: python -m unittest discover -s tests -p 'test_*.py' -v\n"
        "• Rule Benchmark: python tests/test_detection_rules.py\n"
        "• Live UI Dashboard: http://localhost:5173/\n"
        "• Default Admin Credentials: admin / admin123"
    )

    doc.save(str(output_path))
    print(f"Generated DOCX: {output_path}")


def build_html(output_path: Path):
    import markdown
    md_content = (PROJECT_ROOT / "SUBMISSION_REPORT.md").read_text(encoding="utf-8")
    html_body = markdown.markdown(md_content, extensions=["tables", "fenced_code"])

    html_full = f"""<!DOCTYPE html>
<html>
<head>
    <meta charset="utf-8">
    <title>Cybersecurity Assessment Submission Report - Hemanth Prabhu</title>
    <style>
        body {{
            font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Helvetica, Arial, sans-serif;
            line-height: 1.6;
            max-width: 900px;
            margin: 40px auto;
            padding: 0 20px;
            color: #24292e;
        }}
        h1, h2, h3 {{ color: #1a365d; border-bottom: 1px solid #eaecef; padding-bottom: 0.3em; }}
        code {{ background-color: #f6f8fa; padding: 0.2em 0.4em; border-radius: 3px; font-family: Consolas, monospace; }}
        pre {{ background-color: #f6f8fa; padding: 16px; border-radius: 6px; overflow: auto; font-family: Consolas, monospace; }}
        table {{ border-collapse: collapse; width: 100%; margin: 16px 0; }}
        th, td {{ border: 1px solid #dfe2e5; padding: 8px 12px; text-align: left; }}
        th {{ background-color: #f0f4f8; font-weight: 600; color: #1a365d; }}
        tr:nth-child(2n) {{ background-color: #f8fafc; }}
        blockquote {{ border-left: 4px solid #0366d6; color: #586069; padding: 0 16px; margin: 0; }}
        @media print {{
            body {{ max-width: 100%; margin: 20mm; }}
            h1, h2 {{ page-break-after: avoid; }}
            table {{ page-break-inside: avoid; }}
        }}
    </style>
</head>
<body>
    {html_body}
</body>
</html>"""
    output_path.write_text(html_full, encoding="utf-8")
    print(f"Generated HTML: {output_path}")


def package_google_drive_folder():
    SUBMISSION_DIR.mkdir(parents=True, exist_ok=True)

    # 1. Generate DOCX and HTML reports
    docx_path = SUBMISSION_DIR / "Cybersecurity_Assessment_Report_Hemanth_Prabhu.docx"
    html_path = SUBMISSION_DIR / "Cybersecurity_Assessment_Report_Hemanth_Prabhu.html"
    build_docx(docx_path)
    build_html(html_path)

    # 2. Copy benchmark reports
    reports_target = SUBMISSION_DIR / "Benchmark_and_Test_Reports"
    reports_target.mkdir(exist_ok=True)
    shutil.copy2(PROJECT_ROOT / "reports" / "detection_rule_benchmark.txt", reports_target)
    shutil.copy2(PROJECT_ROOT / "reports" / "unit_test_results.txt", reports_target)

    # 3. Copy README and Markdown report
    shutil.copy2(PROJECT_ROOT / "SUBMISSION_REPORT.md", SUBMISSION_DIR)
    shutil.copy2(PROJECT_ROOT / "README.md", SUBMISSION_DIR)

    # 4. Create clean Source Code ZIP
    zip_dest = SUBMISSION_DIR / "OpenSentinel-Lite-SourceCode.zip"
    print("Creating clean source code zip...")
    with zipfile.ZipFile(zip_dest, "w", zipfile.ZIP_DEFLATED) as zf:
        for root, dirs, files in os.walk(PROJECT_ROOT):
            # Exclude virtual envs, node_modules, git, and build artifacts
            dirs[:] = [d for d in dirs if d not in {"venv", ".venv", "node_modules", "dist", "__pycache__", ".git"}]
            for file in files:
                if file.endswith((".pyc", ".db", ".log")) or file == "OpenSentinel-Lite-SourceCode.zip":
                    continue
                file_path = Path(root) / file
                rel_path = file_path.relative_to(PROJECT_ROOT)
                zf.write(file_path, rel_path)

    # 5. Create Submission_Info.txt
    info_text = """CYBERSECURITY ASSESSMENT SUBMISSION
=====================================================
Student Name: Hemanth Prabhu
Student Email: hemanthprabhu2003@gmail.com
Submission Date: October 29, 2026

Selected Assessment Tasks:
1. Task 13: Build a Small Defensive Security Tool (OpenSentinel Lite)
2. Task 9: Test a Defensive Detection Rule

Submission Links:
- GitHub Repository Link: https://github.com/happygeek21/OpenSentinel-Lite-Development
- Live UI Dashboard: http://localhost:5173/
- API Swagger Documentation: http://127.0.0.1:8000/docs
- Default Credentials: admin / admin123

Files included in this folder:
1. Cybersecurity_Assessment_Report_Hemanth_Prabhu.docx (Full Word Report)
2. Cybersecurity_Assessment_Report_Hemanth_Prabhu.html (Printable / PDF-ready HTML)
3. SUBMISSION_REPORT.md (GitHub Flavored Markdown Report)
4. Benchmark_and_Test_Reports/ (Automated test logs and benchmark metrics)
5. OpenSentinel-Lite-SourceCode.zip (Clean repository source code)
=====================================================
"""
    (SUBMISSION_DIR / "Submission_Info.txt").write_text(info_text, encoding="utf-8")
    print(f"Successfully packaged Google Drive folder: {SUBMISSION_DIR}")


if __name__ == "__main__":
    package_google_drive_folder()

